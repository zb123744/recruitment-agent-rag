from __future__ import annotations

from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from sklearn.cluster import KMeans
from sklearn.decomposition import PCA
from sklearn.metrics import silhouette_score
from sklearn.preprocessing import StandardScaler

ROOT = Path(__file__).resolve().parents[2]
DATA = ROOT / "data" / "cleaned" / "recruitment_jobs_cleaned.csv"
FIG = ROOT / "outputs" / "figures"
TAB = ROOT / "outputs" / "tables"
REP = ROOT / "outputs" / "reports"
FIG.mkdir(parents=True, exist_ok=True); TAB.mkdir(parents=True, exist_ok=True); REP.mkdir(parents=True, exist_ok=True)

plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "DejaVu Sans"]
plt.rcParams["axes.unicode_minus"] = False

def main() -> None:
    df = pd.read_csv(DATA, encoding="utf-8-sig")
    features = ["salary_avg_k", "benefit_count", "description_length", "skill_count", "quality_score"]
    X = df[features].apply(pd.to_numeric, errors="coerce")
    X = X.fillna(X.median()).replace([np.inf, -np.inf], np.nan).fillna(0)
    df[features] = X
    scaler = StandardScaler(); Z = scaler.fit_transform(X)
    ks = list(range(2, 9)); inertias=[]; silhouettes=[]
    for k in ks:
        model = KMeans(n_clusters=k, random_state=42, n_init=20).fit(Z)
        inertias.append(model.inertia_); silhouettes.append(silhouette_score(Z, model.labels_))
    best_k = ks[int(np.argmax(silhouettes))]
    model = KMeans(n_clusters=best_k, random_state=42, n_init=20).fit(Z)
    df["cluster"] = model.labels_ + 1
    pca = PCA(n_components=2, random_state=42); coords = pca.fit_transform(Z)
    df["pca1"], df["pca2"] = coords[:,0], coords[:,1]
    pd.DataFrame({"k":ks,"inertia":inertias,"silhouette":silhouettes}).to_csv(TAB/"clustering_k_selection.csv", index=False, encoding="utf-8-sig")
    # Use group size for the record count because source identifiers may be missing.
    profile = df.groupby("cluster").agg(岗位数=("cluster","size"), 平均月薪=("salary_avg_k","mean"), 平均福利数=("benefit_count","mean"), 平均描述长度=("description_length","mean"), 平均技能数=("skill_count","mean"), 平均质量分=("quality_score","mean")).round(2)
    profile.to_csv(TAB/"cluster_profiles.csv", encoding="utf-8-sig")
    # Selection diagnostic
    fig, ax1 = plt.subplots(figsize=(8,4.8)); ax1.plot(ks,inertias,"o-",color="#1e6b55",label="肘部指标"); ax1.set_xlabel("聚类数 K"); ax1.set_ylabel("簇内平方和",color="#1e6b55"); ax2=ax1.twinx(); ax2.plot(ks,silhouettes,"s--",color="#c9993c",label="轮廓系数"); ax2.set_ylabel("轮廓系数",color="#c9993c"); ax1.axvline(best_k,color="#397a9e",alpha=.5); fig.tight_layout(); fig.savefig(FIG/"09_cluster_k_selection.png",dpi=180); plt.close(fig)
    # PCA scatter
    plt.figure(figsize=(8,5.2)); colors=["#1e6b55","#397a9e","#c9993c","#b75d52","#7956a8","#5d8b8b","#cc7f3b","#6b7280"]
    for i in range(1,best_k+1):
        part=df[df.cluster==i]; plt.scatter(part.pca1,part.pca2,s=8,alpha=.35,label=f"类别 {i}",c=colors[(i-1)%len(colors)])
    plt.xlabel(f"主成分1（{pca.explained_variance_ratio_[0]*100:.1f}%）"); plt.ylabel(f"主成分2（{pca.explained_variance_ratio_[1]*100:.1f}%）"); plt.title("岗位特征 PCA 聚类分布"); plt.legend(frameon=False,ncol=2); plt.tight_layout(); plt.savefig(FIG/"10_cluster_pca.png",dpi=180); plt.close()
    # profile chart
    plot_profile=profile[["平均月薪","平均福利数","平均技能数","平均质量分"]].copy(); plot_profile=(plot_profile-plot_profile.min())/(plot_profile.max()-plot_profile.min()).replace(0,1)
    ax=plot_profile.plot(kind="bar",figsize=(9,5),color=["#1e6b55","#70b89e","#c9993c","#397a9e"]); ax.set_xlabel("聚类类别"); ax.set_ylabel("归一化均值"); ax.set_title("各聚类岗位画像对比"); ax.legend(title="指标",frameon=False); plt.xticks(rotation=0); plt.tight_layout(); plt.savefig(FIG/"11_cluster_profiles.png",dpi=180); plt.close()
    make_doc(df, profile, best_k, max(silhouettes), pca)

def make_doc(df, profile, best_k, best_sil, pca):
    doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.7); sec.left_margin=Inches(.8); sec.right_margin=Inches(.8)
    styles=doc.styles; styles["Normal"].font.name="Microsoft YaHei"; styles["Normal"]._element.rPr.rFonts.set(__import__('docx').oxml.ns.qn('w:eastAsia'),'微软雅黑'); styles["Normal"].font.size=Pt(10.5)
    title=doc.add_paragraph(style="Title"); title.alignment=WD_ALIGN_PARAGRAPH.CENTER; title.add_run("多模型对比分析报告")
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run("招聘就业岗位数据的聚类、评分、匹配与分类模型分析").italic=True
    doc.add_heading("一 分析目标与数据", level=1); doc.add_paragraph(f"本报告基于清洗后的招聘岗位数据，共 {len(df):,} 条记录。目标是从岗位画像、岗位质量、简历匹配和质量标签预测四个角度形成可解释的模型对比，为招聘就业分析提供岗位分层、推荐排序和质量识别依据。")
    doc.add_heading("二 K-Means 聚类实现", level=1); doc.add_paragraph("聚类输入包括平均月薪、福利数量、岗位描述长度、技能数量和质量分。先用中位数补齐数值缺失，再进行标准化，避免量纲差异主导距离计算。K 值在 2 至 8 之间搜索，分别计算肘部指标和轮廓系数。")
    doc.add_paragraph(f"综合轮廓系数选择 K={best_k}，最高轮廓系数为 {best_sil:.3f}。PCA 前两个主成分解释 {pca.explained_variance_ratio_.sum()*100:.1f}% 的标准化特征方差，用于二维展示而非替代聚类计算。")
    doc.add_picture(str(FIG/"09_cluster_k_selection.png"), width=Inches(6.4)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER; doc.add_paragraph("图1 K 值选择诊断图", style="Caption").alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_picture(str(FIG/"10_cluster_pca.png"), width=Inches(6.4)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER; doc.add_paragraph("图2 PCA 聚类分布图", style="Caption").alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_heading("三 聚类类别命名与业务解读", level=1); doc.add_paragraph("类别命名依据各簇在薪资、技能、福利、描述完整度和质量分上的相对均值，属于面向业务的解释性命名，不代表岗位原始职类。")
    names={}
    for c,row in profile.iterrows():
        if row["平均质量分"]==profile["平均质量分"].max(): names[c]="高质量综合岗位"
        elif row["平均月薪"]==profile["平均月薪"].max(): names[c]="高薪技能岗位"
        elif row["平均描述长度"]==profile["平均描述长度"].min(): names[c]="低信息基础岗位"
        else: names[c]="中等质量成长岗位"
    table=doc.add_table(rows=1, cols=7); table.style="Table Grid"; hdr=table.rows[0].cells
    for i,h in enumerate(["类别","业务命名","岗位数","平均月薪(千元)","平均福利数","平均技能数","平均质量分"]): hdr[i].text=h
    for c,row in profile.iterrows():
        cells=table.add_row().cells; vals=[f"类别 {c}",names[c],int(row.岗位数),row.平均月薪,row.平均福利数,row.平均技能数,row.平均质量分]
        for i,v in enumerate(vals): cells[i].text=str(v)
    doc.add_picture(str(FIG/"11_cluster_profiles.png"), width=Inches(6.4)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER; doc.add_paragraph("图3 聚类岗位画像对比", style="Caption").alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("业务洞察：高质量综合岗位可作为重点推荐池；高薪技能岗位适合面向具备明确技术栈的求职者做定向推荐；低信息基础岗位需要在数据展示中标注信息完整度风险，避免仅凭薪资进行排序。")
    doc.add_heading("四 评分模型、匹配模型与分类模型对比", level=1)
    doc.add_paragraph("评分模型用薪资、福利、描述完整度和技能数量构造岗位质量分，并以中位数划分高质量标签；人岗匹配模型以简历文本和岗位文本为输入，使用 TF-IDF 与余弦相似度计算技能相似度，同时融合经验、学历、地点和职责词重合度；分类模型则学习质量标签，输出岗位质量预测。三者的业务角色分别是排序、推荐和自动判别。")
    metrics=pd.read_csv(TAB/"model_metrics.csv")
    t=doc.add_table(rows=1, cols=6); t.style="Table Grid"; hs=["模型","准确率","精确率","召回率","F1","ROC-AUC"]
    for i,h in enumerate(hs): t.rows[0].cells[i].text=h
    for _,r in metrics.iterrows():
        cells=t.add_row().cells; vals=[r.model]+[f"{r[k]*100:.2f}%" for k in ["accuracy","precision","recall","f1","roc_auc"]]
        for i,v in enumerate(vals): cells[i].text=str(v)
    doc.add_paragraph("对比结论：XGBoost 与随机森林的 F1 接近 1，适合批量质量识别；SVM 在非线性边界下表现稳定；逻辑回归指标略低但解释性和部署成本较好。实际落地时建议以 XGBoost 作为主模型、逻辑回归作为可解释基线，并保留聚类结果用于岗位分群展示。")
    doc.add_heading("五 局限与优化方向", level=1); doc.add_paragraph("当前聚类特征主要来自结构化字段，尚未将完整岗位文本的语义向量纳入距离计算；分类标签由规则评分生成，存在标签与特征同源导致指标偏高的可能。后续可引入文本 Embedding、时间切分验证、人工抽样复核标签，并对不同来源分别建模，增强跨平台泛化能力。")
    doc.add_heading("六 复现实验文件", level=1); doc.add_paragraph("聚类脚本：src/clustering/run_clustering_report.py；K 值评估表：outputs/tables/clustering_k_selection.csv；聚类画像表：outputs/tables/cluster_profiles.csv；图1至图3位于 outputs/figures。")
    out=REP/"多模型对比分析报告.docx"; doc.save(out)

if __name__ == "__main__": main()
