from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from matching import match_resume

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'outputs' / 'reports' / '人岗匹配算法设计与实现文档.docx'

def shade(cell, fill):
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); cell._tc.get_or_add_tcPr().append(shd)

def set_cell(cell, text, bold=False, size=9, color='000000'):
    cell.text = str(text); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for p in cell.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
            r.bold = bold; r.font.name = 'Microsoft YaHei'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'微软雅黑'); r.font.size = Pt(size); r.font.color.rgb = RGBColor.from_string(color)

def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(headers): set_cell(t.rows[0].cells[i], h, True, 9, 'FFFFFF'); shade(t.rows[0].cells[i], '2F5597')
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,v in enumerate(row):
            set_cell(cells[i], v, False, 9)
            if ri%2: shade(cells[i], 'F2F6FC')
    doc.add_paragraph(); return t

def para(doc, text, style=None):
    p=doc.add_paragraph(style=style); p.paragraph_format.space_after=Pt(6); p.add_run(text); return p

def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    d=Document(); sec=d.sections[0]; sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.7); sec.left_margin=Inches(.8); sec.right_margin=Inches(.8)
    normal=d.styles['Normal']; normal.font.name='Microsoft YaHei'; normal._element.rPr.rFonts.set(qn('w:eastAsia'),'微软雅黑'); normal.font.size=Pt(10.5)
    for name,size in [('Title',20),('Heading 1',15),('Heading 2',12)]:
        s=d.styles[name]; s.font.name='Microsoft YaHei'; s._element.rPr.rFonts.set(qn('w:eastAsia'),'微软雅黑'); s.font.size=Pt(size); s.font.color.rgb=RGBColor(0,0,0)
    p=d.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('人岗匹配算法设计与实现文档')
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('招聘就业岗位数据项目').italic=True
    para(d,'本文档说明招聘就业岗位项目的人岗匹配算法。系统输入候选人简历和岗位信息，输出综合匹配分、匹配等级、已匹配技能、缺失技能及各维度分项结果，为岗位筛选和简历优化提供可解释依据。')
    d.add_heading('一 业务目标与使用场景',1)
    para(d,'面向应届生和招聘人员，解决“简历是否满足岗位要求”“哪些技能已经匹配”“还缺哪些技能”三个问题。算法用于辅助排序，不替代人工面试，也不构成录用承诺。')
    d.add_heading('二 输入处理设计',1)
    table(d,['输入类型','处理方式','输出'],[
        ['粘贴简历','接收纯文本，统一空白字符和标点，保留教育、年限、技能、项目和地点信息','resume_text'],
        ['PDF简历','使用 pypdf.PdfReader 逐页提取文字并拼接；提取后进入同一清洗流程','resume_text'],
        ['岗位信息','将岗位名称、职责、技能、学历、经验、地点等字段拼接成岗位文本','job_text'],
    ])
    para(d,'PDF扫描件如果没有文本层，提取结果可能为空；此类文件需要先OCR，再调用匹配模块。')
    d.add_heading('三 匹配维度与权重',1)
    table(d,['维度','权重','计算方法','业务解释'],[
        ['技能相似度','50%','岗位技能与简历技能的TF-IDF余弦相似度','核心能力匹配程度'],
        ['工作经验','20%','实际年限/要求年限，最高1分','是否达到岗位年限要求'],
        ['学历','15%','博士、硕士、本科、大专、高中分级比较','教育背景是否满足门槛'],
        ['地点','10%','城市文本一致得1分，否则0分；任一为空按1分处理','通勤或工作地适配性'],
        ['职责关键词','5%','简历与岗位共同关键词数量归一化','职责语义相关程度'],
    ])
    d.add_heading('四 核心算法与评分公式',1)
    para(d,'首先从简历和岗位文本中识别技能词表，包括 Python、Java、SQL、Excel、Pandas、Spark、机器学习、深度学习等。技能文本使用 TF-IDF 向量化，再计算余弦相似度：')
    para(d,'TF-IDF(t,d) = TF(t,d) × log(N / (DF(t)+1))')
    para(d,'similarity(A,B) = (A·B) / (||A|| × ||B||)')
    para(d,'最终匹配分：总分 = 100 × (0.50×技能相似度 + 0.20×经验分 + 0.15×学历分 + 0.10×地点分 + 0.05×职责分)。各分项均归一化到0至1。')
    table(d,['分数区间','匹配等级','建议'],[['80–100','高度匹配','优先推荐'],['60–79','较为匹配','补充核验关键条件'],['40–59','部分匹配','针对缺失技能优化简历'],['0–39','匹配度较低','不建议作为首选']])
    d.add_heading('五 实现流程',1)
    for x in ['读取粘贴文本或PDF文本。','识别技能、学历、工作年限和地点等结构化信息。','使用TF-IDF和余弦相似度计算技能分。','分别计算经验、学历、地点和职责关键词分。','按权重加权得到总分并映射匹配等级。','输出匹配技能、缺失技能、分项得分和解释。']:
        para(d,x,'List Number')
    d.add_heading('六 代码实现说明',1)
    para(d,'可运行代码位于 src/matching/matching.py，主要接口为 match_resume(resume_text, job_text, resume_location="", job_location="")。extract_pdf(path) 负责PDF文本提取，match_resume 负责完整评分。技能词表、学历等级和经验年限解析规则均在模块中集中维护，便于后续扩展。')
    d.add_heading('七 示例计算',1)
    resume='本科，3年经验。熟悉 Python、SQL、Pandas、机器学习。项目地点上海。'
    job='本科，要求3年经验，熟悉Python、SQL、Pandas、Spark、机器学习。工作地点上海。'
    r=match_resume(resume,job,'上海','上海')
    para(d,'示例简历：本科，3年经验，Python、SQL、Pandas、机器学习，地点上海。\n示例岗位：本科，要求3年经验，Python、SQL、Pandas、Spark、机器学习，地点上海。')
    table(d,['指标','结果'],[['综合匹配分',f"{r['score']}分"],['匹配等级',r['level']],['技能相似度',r['skill_similarity']],['已匹配技能','、'.join(r['skill_overlap'])],['缺失技能','、'.join(r['missing_skills']) or '无'],['经验/学历/地点/职责分',f"{r['experience_score']} / {r['education_score']} / {r['location_score']} / {r['responsibility_score']}"]])
    d.add_heading('八 结果解释与应用',1)
    para(d,'系统不仅输出一个分数，还输出缺失技能。例如示例岗位缺少 Spark 时，候选人可以据此补充项目经历或学习计划；招聘人员可先按综合分排序，再结合企业硬性条件、项目质量和面试结果做最终判断。')
    d.add_heading('九 局限性与优化方向',1)
    table(d,['问题','影响','优化方向'],[
        ['TF-IDF依赖词面重合','难识别同义词和上下文能力','引入中文句向量Embedding并与关键词分融合'],
        ['地点、学历表达不统一','可能出现误判','建立城市别名、学历规范化和硬条件过滤'],
        ['扫描PDF无文本层','简历内容无法进入匹配','增加OCR识别和人工校验'],
        ['技能词表有限','新技术或缩写可能漏识别','从岗位数据统计高频技能并持续更新词典'],
        ['权重来自业务设定','不同岗位重要性不同','按岗位类别学习或校准权重'],
    ])
    d.add_heading('十 交付文件',1)
    table(d,['文件','用途'],[['src/matching/matching.py','人岗匹配核心算法与PDF提取接口'],['docs/04_人岗匹配文档.md','Markdown版设计说明'],['outputs/reports/人岗匹配算法设计与实现文档.docx','提交用Word文档']])
    d.save(OUT); print(OUT)

if __name__=='__main__': main()
